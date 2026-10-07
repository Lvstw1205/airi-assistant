"""Bounded local document text extraction; never executes macros or attachments."""
import csv, io, json, re, resource, subprocess, sys, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

LIMIT=60000

def extract(path):
    suffix=path.suffix.lower(); parts=[]; size=0; truncated=False
    def add(value):
        nonlocal size,truncated
        value=str(value).strip()
        if not value:return
        if size+len(value)+1>LIMIT:truncated=True
        if size<LIMIT:
            parts.append(value[:LIMIT-size]);size+=len(parts[-1])+1
    if suffix in ('.xlsx','.docx','.pptx'):
        with zipfile.ZipFile(path) as archive:
            items=archive.infolist()
            if len(items)>10000 or sum(x.file_size for x in items)>120*1024*1024:
                raise ValueError('압축을 풀었을 때 너무 큰 문서입니다. 파일을 나눠 주세요.')
    if suffix in ('.md','.markdown','.txt','.csv','.tsv','.json'):
        raw=path.read_bytes();text=None
        for encoding in ('utf-8-sig','utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'cp949'):
            try:text=raw.decode(encoding);break
            except UnicodeError:pass
        if text is None or '\0' in text:raise ValueError('UTF-8 또는 한국어 텍스트 파일로 저장해 주세요.')
        add(text)
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('암호를 해제한 PDF를 첨부해 주세요.')
        for i,page in enumerate(reader.pages):
            if size>=LIMIT:truncated=True;break
            add('[페이지 %d]\n%s'%(i+1,page.extract_text() or ''))
        if not any(re.search(r'[^\s\d\[\]페이지]',p) for p in parts):
            raise ValueError('읽을 수 있는 텍스트가 없는 PDF입니다. 스캔 파일은 OCR 후 첨부해 주세요.')
    elif suffix=='.xlsx':
        import openpyxl
        book=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                add('[시트: '+sheet.title+']')
                for i,row in enumerate(sheet.iter_rows(values_only=True)):
                    if size>=LIMIT or i>=10000:truncated=True;break
                    if len(row)>200:truncated=True
                    add('\t'.join('' if v is None else str(v) for v in row[:200]))
        finally:book.close()
    elif suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(path,on_demand=True)
        try:
            for sheet in book.sheets():
                add('[시트: '+sheet.name+']')
                for i in range(min(sheet.nrows,10000)):
                    if size>=LIMIT:truncated=True;break
                    add('\t'.join(str(v) for v in sheet.row_values(i)[:200]))
                if sheet.nrows>10000 or sheet.ncols>200:truncated=True
        finally:book.release_resources()
    elif suffix=='.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for block in Document(path).iter_inner_content():
            if size>=LIMIT:truncated=True;break
            if isinstance(block,Paragraph):add(block.text)
            elif isinstance(block,Table):
                for row in block.rows:add('\t'.join(c.text for c in row.cells))
    elif suffix=='.pptx':
        with zipfile.ZipFile(path) as archive:
            slides=sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
            for i,name in enumerate(slides):
                if size>=LIMIT:truncated=True;break
                root=ET.fromstring(archive.read(name));add('[슬라이드 %d]\n%s'%(i+1,'\n'.join(n.text or '' for n in root.iter() if n.tag.endswith('}t'))))
    elif suffix in ('.doc','.rtf'):
        converted=subprocess.run(['/usr/bin/textutil','-convert','txt','-stdout',str(path)],capture_output=True,timeout=25)
        if converted.returncode:raise ValueError('문서를 읽지 못했습니다. DOCX로 다시 저장해 주세요.')
        add(converted.stdout.decode('utf-8'))
    else:raise ValueError('지원하지 않는 파일 형식입니다.')
    text='\n'.join(parts)
    if not text.strip():raise ValueError('파일에 읽을 수 있는 내용이 없습니다.')
    return {'text':text,'truncated':truncated,'chars':len(text)}

if __name__=='__main__':
    resource.setrlimit(resource.RLIMIT_CPU,(35,35))
    try:print(json.dumps(extract(Path(sys.argv[1])),ensure_ascii=False))
    except Exception as e:
        print(json.dumps({'error':str(e) if isinstance(e,ValueError) else '손상되었거나 지원되지 않는 문서입니다. 다른 형식으로 저장해 주세요.'},ensure_ascii=False));sys.exit(1)
